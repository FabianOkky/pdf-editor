"""Tests for the smart Word export endpoint (``/pdf/export/docx``).

The pipeline must: convert a native PDF straight through pdf2docx; route a scanned PDF through
OCR *first* so its text is recoverable; and report which path ran. The native path and the
routing decision are tested without the OCR engine (the engine is faked), so they always run;
the real scanned → ``.docx`` proof is skipped when tessdata is absent. Inputs are never mutated.
"""

from __future__ import annotations

import base64
import io
import zipfile

import fitz
import pytest
from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT

from app.services import pdf_export
from app.services.pdf_ocr import ocr_available

requires_ocr = pytest.mark.skipif(
    not ocr_available(),
    reason="OCR language data (tessdata) not installed; see pdf-service/README.md.",
)


def _export(client, secret, pdf: bytes, language: str | None = None):
    data = {"language": language} if language is not None else None
    return client.post(
        "/pdf/export/docx",
        headers={"X-Pdf-Secret": secret},
        files={"file": ("doc.pdf", pdf, "application/pdf")},
        data=data,
    )


def _docx_text(raw: bytes) -> str:
    """Extract the visible paragraph text from a ``.docx`` byte string."""
    document = DocxDocument(io.BytesIO(raw))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _one_page_pdf(text: str) -> bytes:
    doc = fitz.open()
    doc.new_page(width=612, height=792).insert_text((72, 72), text, fontsize=14)
    data = doc.tobytes()
    doc.close()
    return data


def test_export_requires_secret(client, native_pdf):
    response = client.post(
        "/pdf/export/docx",
        files={"file": ("doc.pdf", native_pdf, "application/pdf")},
    )
    assert response.status_code == 401


def test_export_rejects_a_corrupt_pdf(client, secret):
    assert _export(client, secret, b"definitely not a pdf").status_code == 422


def test_export_native_pdf_produces_editable_docx(client, secret, native_pdf):
    response = _export(client, secret, native_pdf)
    assert response.status_code == 200

    body = response.json()
    assert body["source_type"] == "native"
    assert body["ocr_applied"] is False
    assert body["page_count"] == 2

    raw = base64.b64decode(body["content_base64"])
    assert zipfile.is_zipfile(io.BytesIO(raw))  # a .docx is an OpenXML zip package
    assert "Hello world" in _docx_text(raw)


def test_export_native_skips_ocr(client, secret, native_pdf, monkeypatch):
    """A native PDF must never hit the (expensive) OCR step."""

    def must_not_run(*args, **kwargs):
        raise AssertionError("OCR must not run for a native PDF")

    monkeypatch.setattr(pdf_export, "ocr_pdf", must_not_run)

    body = _export(client, secret, native_pdf).json()
    assert body["ocr_applied"] is False


def test_export_routes_scanned_through_ocr_first(client, secret, scanned_pdf, monkeypatch):
    """A scanned PDF is OCR'd first; the OCR output (not the original) feeds pdf2docx."""
    calls = {"count": 0}

    def fake_ocr(data, language=None, tessdata=None, cache_dir=None):
        calls["count"] += 1
        searchable = _one_page_pdf("FAKEOCR sentinel 4815")
        return {
            "page_count": 1,
            "content_base64": base64.b64encode(searchable).decode("ascii"),
            "text": "FAKEOCR sentinel 4815",
            "language": "eng",
        }

    monkeypatch.setattr(pdf_export, "ocr_pdf", fake_ocr)

    body = _export(client, secret, scanned_pdf).json()
    assert body["source_type"] == "scanned"
    assert body["ocr_applied"] is True
    assert calls["count"] == 1
    # The faked OCR text flowed into the produced .docx, proving OCR-first wiring.
    assert "FAKEOCR" in _docx_text(base64.b64decode(body["content_base64"]))


@requires_ocr
def test_export_scanned_pdf_recovers_text_with_real_ocr(client, secret, scanned_text_pdf):
    response = _export(client, secret, scanned_text_pdf)
    assert response.status_code == 200

    body = response.json()
    assert body["source_type"] == "scanned"
    assert body["ocr_applied"] is True

    text = _docx_text(base64.b64decode(body["content_base64"]))
    # The whole point: a scanned doc that naive converters leave empty comes back with text.
    assert "4815" in text


def _empty_docx_bytes() -> bytes:
    document = DocxDocument()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_recover_dropped_text_appends_text_missing_from_the_docx():
    """When the DOCX lacks source text, the recovery step appends it (no content lost)."""
    pdf = _one_page_pdf("Margomulyo sentinel 73519")

    recovered = pdf_export._recover_dropped_text(pdf, _empty_docx_bytes())

    text = _docx_text(recovered)
    assert "Recovered text" in text
    assert "Margomulyo sentinel 73519" in text


def test_recover_dropped_text_recovers_partial_line_drops():
    """Even a few words dropped from a mostly-present line are recovered (not just whole lines)."""
    pdf = _one_page_pdf("Alpha Bravo Charlie Delta Echo")

    document = DocxDocument()
    document.add_paragraph("Alpha Bravo Delta")  # "Charlie" and "Echo" were dropped
    buffer = io.BytesIO()
    document.save(buffer)

    recovered = pdf_export._recover_dropped_text(pdf, buffer.getvalue())

    text = _docx_text(recovered)
    assert "Recovered text" in text
    assert "Charlie" in text
    assert "Echo" in text


def test_recover_dropped_text_is_a_noop_when_all_text_is_present():
    """When every source word already appears in the DOCX, the bytes are returned unchanged."""
    pdf = _one_page_pdf("Margomulyo sentinel 73519")

    document = DocxDocument()
    document.add_paragraph("Margomulyo sentinel 73519")
    buffer = io.BytesIO()
    document.save(buffer)
    docx_bytes = buffer.getvalue()

    assert pdf_export._recover_dropped_text(pdf, docx_bytes) == docx_bytes


def test_layout_export_carries_page_geometry_and_font_size():
    """The rebuilt (OCR/fallback) document keeps the PDF's page size and per-block font size."""
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)  # A4 landscape, in points
    page.insert_text((60, 80), "Quarterly Report", fontsize=24)
    page.insert_text((60, 140), "Body copy at a normal reading size.", fontsize=11)
    pdf = doc.tobytes()
    doc.close()

    raw, page_count = pdf_export._layout_to_docx(pdf)
    document = DocxDocument(io.BytesIO(raw))

    assert page_count == 1

    section = document.sections[0]
    assert section.page_width.pt == pytest.approx(842, abs=1)
    assert section.page_height.pt == pytest.approx(595, abs=1)
    assert section.orientation == WD_ORIENT.LANDSCAPE

    sizes = {
        paragraph.text: paragraph.runs[0].font.size.pt
        for paragraph in document.paragraphs
        if paragraph.runs
    }
    assert sizes["Quarterly Report"] == pytest.approx(24, abs=1)
    assert sizes["Body copy at a normal reading size."] == pytest.approx(11, abs=1)


def test_layout_export_emits_larger_text_in_bold_as_a_heading():
    """A span well above the body size reads as a heading, so it comes through bold."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 90), "Section Title", fontsize=22)
    for offset in range(4):
        page.insert_text((72, 140 + offset * 20), f"Ordinary body line {offset}.", fontsize=10)
    pdf = doc.tobytes()
    doc.close()

    document = DocxDocument(io.BytesIO(pdf_export._layout_to_docx(pdf)[0]))
    bold_by_text = {
        paragraph.text: paragraph.runs[0].bold
        for paragraph in document.paragraphs
        if paragraph.runs
    }

    assert bold_by_text["Section Title"] is True
    assert bold_by_text["Ordinary body line 0."] is False


def _block(x0: float, y0: float, x1: float, y1: float, text: str) -> dict:
    """A minimal layout block, as :func:`pdf_export._page_layout` would produce it."""
    return {"bbox": (x0, y0, x1, y1), "lines": [text], "size": 10.0, "bold": False, "italic": False}


def test_reading_order_reads_two_column_pages_column_by_column():
    """When a page splits cleanly at the gutter, the left column is read out in full first."""
    blocks = [_block(60, 100 + row * 40, 280, 120 + row * 40, f"left{row}") for row in range(4)]
    blocks += [_block(330, 100 + row * 40, 550, 120 + row * 40, f"right{row}") for row in range(4)]

    ordered = [block["lines"][0] for block in pdf_export._reading_order(blocks, page_width=612)]

    assert ordered == ["left0", "left1", "left2", "left3", "right0", "right1", "right2", "right3"]


def test_reading_order_stays_top_to_bottom_when_a_block_spans_the_gutter():
    """A full-width block (a title, a wide table) means the page is not really two columns."""
    blocks = [
        _block(60, 60, 550, 90, "Full width title"),
        _block(60, 120, 280, 140, "left"),
        _block(330, 120, 550, 140, "right"),
    ]

    ordered = [block["lines"][0] for block in pdf_export._reading_order(blocks, page_width=612)]

    assert ordered == ["Full width title", "left", "right"]


def test_word_comparison_folds_ligatures_and_accents():
    """A PDF ligature ("oﬃce") must not look like a word Word failed to write."""
    present = set(pdf_export._words("The office final draft resume"))

    assert not pdf_export._line_is_missing("The oﬃce ﬁnal draft", present)
    assert not pdf_export._line_is_missing("résumé", present)
    assert pdf_export._line_is_missing("genuinely absent wording", present)


def test_mixed_pdf_keeps_native_layout_and_ocrs_only_the_scanned_pages(
    client, secret, mixed_pdf, monkeypatch
):
    """Mixed files convert natively (layout kept) and OCR runs on the image-only page alone."""
    seen = {}

    def fake_ocr(data, language=None, tessdata=None, cache_dir=None):
        with fitz.open("pdf", data) as handed:
            seen["page_count"] = handed.page_count
        searchable = _one_page_pdf("SCANSENTINEL 90210")
        return {
            "page_count": 1,
            "content_base64": base64.b64encode(searchable).decode("ascii"),
            "text": "SCANSENTINEL 90210",
            "language": "eng",
        }

    monkeypatch.setattr(pdf_export, "ocr_pdf", fake_ocr)

    body = _export(client, secret, mixed_pdf).json()

    assert body["source_type"] == "mixed"
    assert body["ocr_applied"] is True
    assert body["page_count"] == 2
    # Only the image-only page was handed to OCR — the native page kept its real layout.
    assert seen["page_count"] == 1

    text = _docx_text(base64.b64decode(body["content_base64"]))
    assert "Hello world" in text  # the native page converted normally
    assert "SCANSENTINEL" in text  # the scanned page's text was folded in


def test_mixed_pdf_still_exports_when_ocr_is_unavailable(client, secret, mixed_pdf, monkeypatch):
    """Missing tessdata degrades to a native-only conversion instead of failing the export."""

    def unavailable(*args, **kwargs):
        raise ValueError("OCR language data 'eng' not found")

    monkeypatch.setattr(pdf_export, "ocr_pdf", unavailable)

    response = _export(client, secret, mixed_pdf)
    assert response.status_code == 200

    body = response.json()
    assert body["ocr_applied"] is False
    assert "Hello world" in _docx_text(base64.b64decode(body["content_base64"]))


def test_native_export_falls_back_to_a_text_layout_when_pdf2docx_fails(
    client, secret, native_pdf, monkeypatch
):
    """A converter blow-up must still yield a readable .docx, never an empty download."""

    def exploding_converter(*args, **kwargs):
        raise RuntimeError("pdf2docx exploded")

    monkeypatch.setattr(pdf_export, "Converter", exploding_converter)

    response = _export(client, secret, native_pdf)
    assert response.status_code == 200

    assert "Hello world" in _docx_text(base64.b64decode(response.json()["content_base64"]))
